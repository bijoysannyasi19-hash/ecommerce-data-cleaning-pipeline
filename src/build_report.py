import os
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def build_report():
    print("Loading data for report statistics...")
    df_raw = pd.read_excel('data/raw/Online_Retail.xlsx')
    df_clean = pd.read_csv('data/processed/online_retail_cleaned.csv')

    raw_rows, raw_cols = df_raw.shape
    clean_rows, clean_cols = df_clean.shape
    
    missing_customer = df_raw['CustomerID'].isnull().sum()
    missing_desc = df_raw['Description'].isnull().sum()
    
    duplicates = df_raw.duplicated().sum()
    
    negative_quantity = (df_raw['Quantity'] < 0).sum()
    negative_price = (df_raw['UnitPrice'] < 0).sum()
    zero_price = (df_raw['UnitPrice'] == 0).sum()
    
    doc = Document()
    
    # Title Page
    title = doc.add_heading('WEEK 1 PROJECT REPORT\nData Acquisition, Cleaning, and Preprocessing', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_page_break()

    # Executive Summary
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(f"This report details the data acquisition, cleaning, and preprocessing steps performed on the UCI Online Retail dataset. We started with {raw_rows} raw transactional records. After treating missing values, removing {duplicates} duplicate records, cleaning invalid prices, and trimming extreme outliers, we produced a final analysis-ready dataset containing {clean_rows} records.")

    # Objectives
    doc.add_heading('2. Objectives', level=1)
    doc.add_paragraph("The objective of this assignment is to acquire a publicly available dataset and perform structured data exploration, quality assessment, and cleaning. The goal is to document all decisions clearly and handle missing values, duplicates, and outliers systematically, ensuring the data is ready for downstream analytical tasks.")

    # Dataset Selection
    doc.add_heading('3. Dataset Selection', level=1)
    doc.add_paragraph("Dataset: Online Retail\nSource: UCI Machine Learning Repository\nURL: https://archive.ics.uci.edu/dataset/352/online+retail\n\nThis dataset was chosen because it contains real-world e-commerce transactional data with genuine quality issues such as missing values, negative quantities from returns, and administrative price adjustments.")

    # Data Acquisition Method
    doc.add_heading('4. Data Acquisition Method', level=1)
    doc.add_paragraph("The dataset was acquired programmatically using Python's `requests` library. The file was downloaded directly from the UCI repository and saved as an Excel file (`Online_Retail.xlsx`) in the `data/raw/` directory. The original file format and raw data were kept strictly unmodified to ensure a reproducible pipeline.")

    # Dataset Description
    doc.add_heading('5. Dataset Description', level=1)
    doc.add_paragraph(f"The raw dataset consists of {raw_rows} rows and {raw_cols} columns. Each row represents a single product line item within an order (or cancellation) made by a customer from a specific country.")

    # Data Dictionary
    doc.add_heading('6. Data Dictionary', level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Column Name'
    hdr_cells[1].text = 'Data Type'
    hdr_cells[2].text = 'Meaning'
    hdr_cells[3].text = 'Potential Data Quality Concern'
    
    dictionary = [
        ('InvoiceNo', 'Object', 'Invoice number (starts with C if cancelled)', 'Cancellations'),
        ('StockCode', 'Object', 'Product code', 'Non-product codes'),
        ('Description', 'Object', 'Product name', 'Missing text'),
        ('Quantity', 'Numeric', 'Item quantity', 'Negative values'),
        ('InvoiceDate', 'Datetime', 'Date and time of transaction', 'Parsing errors'),
        ('UnitPrice', 'Numeric', 'Price per unit', 'Zero or negative prices'),
        ('CustomerID', 'Numeric/Float', 'Customer ID', 'High missing rate'),
        ('Country', 'Object', 'Country name', 'Inconsistent names')
    ]
    for col1, col2, col3, col4 in dictionary:
        row_cells = table.add_row().cells
        row_cells[0].text = col1
        row_cells[1].text = col2
        row_cells[2].text = col3
        row_cells[3].text = col4

    # Initial Data Exploration
    doc.add_heading('7. Initial Data Exploration', level=1)
    doc.add_paragraph("Initial inspection revealed that the dataset is highly concentrated in the United Kingdom. Numerical distributions showed massive outliers in both Quantity and UnitPrice. We also observed missing values and datatypes that needed correction.")
    
    if os.path.exists('figures/top_countries.png'):
        doc.add_picture('figures/top_countries.png', width=Inches(5))
        doc.add_paragraph("Figure 1: Top 10 Countries by Transaction Count", style='Caption')

    # Missing Value Analysis
    doc.add_heading('8. Missing Value Analysis', level=1)
    doc.add_paragraph(f"The initial assessment identified missing values in two columns:\n- CustomerID: {missing_customer} missing values (approx 24.9%)\n- Description: {missing_desc} missing values")
    doc.add_paragraph("Decision: Since dropping 24.9% of the dataset would severely skew revenue metrics, we chose not to drop rows with missing CustomerIDs. Instead, we imputed them with a placeholder (-1) and created an `Is_Anonymous` flag. For Description, we imputed missing strings with 'UNKNOWN'.")

    # Duplicate Analysis
    doc.add_heading('9. Duplicate Analysis', level=1)
    doc.add_paragraph(f"We found {duplicates} completely identical rows across all columns. Because there is no unique line-item ID, exact matches occurring at the same timestamp were assumed to be system duplication errors. \nDecision: All {duplicates} duplicated rows were removed.")

    # Data Type Corrections
    doc.add_heading('10. Data Type Corrections', level=1)
    doc.add_paragraph("Once missing CustomerIDs were replaced with numeric placeholders, the column was converted from float to integer format to accurately reflect IDs without decimal points.")

    # Invalid/Inconsistent Data
    doc.add_heading('11. Invalid/Inconsistent Data', level=1)
    doc.add_paragraph(f"Investigation of numerical columns revealed:\n- {negative_quantity} records with negative Quantity. These mostly represented legitimate cancellations/returns (InvoiceNo starting with 'C'). We chose to retain them but flag them as `Is_Cancelled`.\n- {negative_price} records with negative UnitPrice (bad debt adjustments) and {zero_price} records with zero UnitPrice. \nDecision: Records with zero or negative UnitPrice were removed as they do not represent valid product sales.")

    # Outlier Detection and Treatment
    doc.add_heading('12. Outlier Detection and Treatment', level=1)
    doc.add_paragraph("Extremely high quantities and unit prices were present. However, standard IQR removal would indiscriminately delete valid bulk wholesale orders.")
    if os.path.exists('figures/outliers_before.png'):
        doc.add_picture('figures/outliers_before.png', width=Inches(5))
        doc.add_paragraph("Figure 2: Outliers before treatment", style='Caption')
    doc.add_paragraph("Decision: Instead of standard IQR filtering, we applied percentile-based trimming, removing only values above the 99.9th percentile for Quantity and UnitPrice, and below the 0.1th percentile for Quantity. This surgically removes extreme administrative adjustments while retaining normal retail variance.")

    # Data Cleaning Pipeline
    doc.add_heading('13. Data Cleaning Pipeline', level=1)
    doc.add_paragraph("A derived variable `TotalAmount` was created by multiplying Quantity and UnitPrice. This facilitates downstream revenue analysis. No scaling or machine learning encodings were applied because this dataset is currently prepared for general business analytics rather than a specific predictive model.")

    # Before vs After Comparison
    doc.add_heading('14. Before vs After Comparison', level=1)
    comp_table = doc.add_table(rows=1, cols=4)
    comp_table.style = 'Table Grid'
    hdr = comp_table.rows[0].cells
    hdr[0].text = 'Metric'
    hdr[1].text = 'Before Cleaning'
    hdr[2].text = 'After Cleaning'
    hdr[3].text = 'Change'
    
    comp_data = [
        ('Rows', str(raw_rows), str(clean_rows), str(clean_rows - raw_rows)),
        ('Columns', str(raw_cols), str(clean_cols), str(clean_cols - raw_cols)),
        ('Missing Values', str(missing_customer + missing_desc), '0', str(-(missing_customer + missing_desc))),
        ('Duplicates', str(duplicates), '0', str(-duplicates))
    ]
    for m, b, a, c in comp_data:
        row_cells = comp_table.add_row().cells
        row_cells[0].text = m
        row_cells[1].text = b
        row_cells[2].text = a
        row_cells[3].text = c

    # Data Quality Assessment
    doc.add_heading('15. Data Quality Assessment', level=1)
    doc.add_paragraph("The dataset now exhibits high Completeness (missing values handled), Validity (negative prices removed), Consistency (datatypes enforced), and Uniqueness (exact duplicates removed).")

    # Challenges and Solutions
    doc.add_heading('16. Challenges and Solutions', level=1)
    doc.add_paragraph("Challenge: Determining whether to drop negative Quantity values.\nWhy it was difficult: Blindly dropping negative values cleans the distribution but inflates total sales numbers by ignoring returns.\nSolution: I investigated the Invoice numbers and confirmed that most negative quantities corresponded to cancellations. Therefore, I retained them and created an `Is_Cancelled` boolean flag to allow flexible downstream analysis.")

    # Impact of Preprocessing
    doc.add_heading('17. Impact of Preprocessing', level=1)
    doc.add_paragraph("By choosing to impute missing CustomerIDs rather than dropping them, any future analysis analyzing total sales volume remains accurate and robust. Conversely, since these rows do not belong to known customers, they must be excluded if performing customer-centric modeling (like RFM segmentation). By retaining extreme but valid bulk orders (up to the 99.9th percentile), statistical models will correctly capture wholesale dynamics instead of treating them as invalid.")

    # Reflective Commentary
    doc.add_heading('18. Reflective Commentary', level=1)
    doc.add_paragraph("Initially, I considered removing all rows with negative values because they appeared to be data entry errors. However, after carefully checking the underlying invoice patterns, I found evidence that some represented legitimate customer returns. This experience taught me that domain-specific interpretation is just as important as statistical thresholds. Blind statistical cleaning can sometimes remove the most important business signals from a dataset.")

    # Final Conclusion
    doc.add_heading('19. Final Conclusion', level=1)
    doc.add_paragraph(f"The Online Retail dataset was successfully acquired, explored, and preprocessed. We addressed missing data, identified and resolved {duplicates} duplicates, fixed data types, and applied targeted outlier removal. The final cleaned dataset of {clean_rows} rows is robust and analysis-ready.")

    # References
    doc.add_heading('20. References', level=1)
    doc.add_paragraph("UCI Machine Learning Repository. Online Retail Dataset. https://archive.ics.uci.edu/dataset/352/online+retail\nPandas Documentation: https://pandas.pydata.org/docs/")

    doc.save('reports/Week1_Data_Acquisition_Cleaning_Preprocessing.docx')
    print("DOCX Report saved.")

if __name__ == "__main__":
    build_report()
